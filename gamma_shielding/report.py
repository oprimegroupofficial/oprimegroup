"""Output writers: Excel workbook, CSV files, PNG plots and a Markdown report (Turkish)."""

from __future__ import annotations

import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

from . import physics as ph  # noqa: E402
from .pipeline import Results  # noqa: E402
from .xsec import CrossSections  # noqa: E402

PALETTE = ["#1f77b4", "#d95f02", "#1b9e77", "#7570b3", "#e7298a", "#66a61e", "#a6761d", "#666666"]


# ------------------------------------------------------------------------------ formatting
def fmt_pm(v, u, k: float = 1.0, nd_min: int = 0) -> str:
    """'value ± U' with U rounded to 2 significant digits (GUM)."""
    if v is None or u is None or not np.isfinite(v):
        return "-"
    if not np.isfinite(u) or u <= 0:
        return f"{v:.4g}"
    U = k * u
    if abs(v) >= 1e6 or (0 < abs(v) < 1e-3):          # scientific notation with a common exponent
        e = int(math.floor(math.log10(abs(v))))
        return f"({fmt_pm(v / 10**e, U / 10**e, 1.0, nd_min)})e{e:+d}".replace("e+", "e")
    dec = max(nd_min, 1 - int(math.floor(math.log10(U)))) if U > 0 else 3
    dec = min(dec, 10)
    return f"{v:.{dec}f} ± {U:.{dec}f}"


# ------------------------------------------------------------------------------ plots
def plot_spectrum(res: Results, name: str, path: Path):
    a = res.analyses[name]
    spec, cal = a.spectrum, a.calibration
    regions = list(a.fits.keys())
    shown = [r for r in regions if any(L.region == r and L.significance >= res.cfg.min_significance
                                       for L in a.lines.values())]
    ncol = max(1, len(shown))
    fig = plt.figure(figsize=(3.6 * ncol, 9.2), constrained_layout=True)
    gs = fig.add_gridspec(3, ncol, height_ratios=[1.5, 2.0, 0.8])
    ax0 = fig.add_subplot(gs[0, :])
    ch = np.arange(spec.n_channels)
    E = cal.energy(ch)
    ax0.semilogy(E, np.clip(spec.counts, 0.8, None), lw=0.6, color="#333333")
    for key, L in a.lines.items():
        if L.significance >= res.cfg.min_significance:
            ax0.axvline(L.E_eff, color="#d95f02", lw=0.6, alpha=0.6)
            ax0.text(L.E_eff, ax0.get_ylim()[1] if False else spec.counts.max() * 0.9, key, rotation=90, va="top",
                     ha="right", fontsize=7, color="#d95f02")
    ax0.set_xlim(0, float(E[int(np.nonzero(spec.counts > 20)[0].max())]) * 1.05)
    ax0.set_xlabel("Enerji (keV)")
    ax0.set_ylabel("Sayım / kanal")
    ax0.set_title(f"{name} - canlı süre {spec.live_time:.0f} s, ölü zaman %{100*spec.dead_time_fraction:.1f}, "
                  f"kalibrasyon rms {cal.rms_residual_keV:.1f} keV", fontsize=10)
    for j, rname in enumerate(shown):
        f = a.fits[rname]["base"]
        ax = fig.add_subplot(gs[1, j])
        xe = cal.energy(f.x)
        ax.step(xe, f.y, where="mid", lw=0.6, color="#333333", label="veri")
        ax.plot(xe, f.model, color="#d62728", lw=1.0, label="model")
        ax.plot(xe, f.background, color="#1f77b4", lw=0.8, ls="--", label="fon")
        for i, (lab, pk) in enumerate(f.peaks.items()):
            ax.plot(xe, f.background + pk["curve"], lw=0.8, color=PALETTE[(i + 2) % len(PALETTE)],
                    label=lab if pk["reported"] else lab + " (yardımcı)")
        ax.set_title(f"{rname}\nχ²/ndf={f.chi2_red:.2f}", fontsize=8)
        ax.tick_params(labelsize=7)
        ax.legend(fontsize=5.5, loc="upper right")
        axr = fig.add_subplot(gs[2, j], sharex=ax)
        axr.plot(xe, (f.y - f.model) / np.sqrt(np.clip(f.model, 1, None)), lw=0.5, color="#333333")
        axr.axhline(0, color="k", lw=0.5)
        axr.set_ylim(-4, 4)
        axr.set_xlabel("keV", fontsize=7)
        axr.tick_params(labelsize=6)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def plot_relative(res: Results, path: Path):
    df = res.relative
    if df.empty:
        return
    fig, ax = plt.subplots(figsize=(8, 4.5), constrained_layout=True)
    samples = list(df["Numune"].unique())
    for i, s in enumerate(samples):
        d = df[df["Numune"] == s].sort_values("E_keV")
        off = 1.0 + 0.03 * (i - (len(samples) - 1) / 2)
        ax.errorbar(d["E_keV"] * off, d["ln(Iref/I)"], yerr=d["u_toplam"], fmt="o", ms=4, capsize=3,
                    color=PALETTE[i], label=s)
    ax.axhline(0, color="k", lw=0.7)
    ax.set_xscale("log")
    ax.set_xlabel("Enerji (keV)")
    ax.set_ylabel(f"ln(I_ref / I)   [ref = {res.status['relative_reference']}]")
    ax.set_title("Göreli zayıflama (bilgi amaçlı; LAC değildir) - hata çubukları: istatistik ⊕ model", fontsize=9)
    ax.legend(fontsize=8)
    ax.grid(alpha=0.3)
    fig.savefig(path, dpi=130)
    plt.close(fig)


def plot_parameters(res: Results, xs: CrossSections, outdir: Path):
    df = res.attenuation
    if df.empty or "LAC_cm-1" not in df.columns or df["LAC_cm-1"].notna().sum() == 0:
        return
    cfg = res.cfg
    egrid = np.geomspace(50, 1500, 60)
    specs = [("LAC_cm-1", "u_LAC_cm-1", "LAC_kuram_cm-1", "Doğrusal zayıflatma katsayısı μ (cm⁻¹)", "lac_enerji.png", True),
             ("MAC_cm2_g", "u_MAC_cm2_g", "MAC_kuram_cm2_g", "Kütle zayıflatma katsayısı μ/ρ (cm²/g)", "mac_enerji.png", True),
             ("HVL_cm", "u_HVL_cm", "HVL_kuram_cm", "Yarı değer kalınlığı HVL (cm)", "hvl_enerji.png", True),
             ("Zeff", "u_Zeff", "Zeff_kuram", "Etkin atom numarası Zeff", "zeff_enerji.png", False),
             ("RPE_yuzde", "u_RPE_yuzde", None, "Radyasyon koruma verimi RPE (%)", "rpe_enerji.png", False)]
    for col, ucol, thcol, label, fname, logy in specs:
        if col not in df.columns or df[col].notna().sum() == 0:
            continue
        fig, ax = plt.subplots(figsize=(8, 5), constrained_layout=True)
        for i, s in enumerate(cfg.samples):
            d = df[(df["Numune"] == s.name) & df[col].notna()].sort_values("E_keV")
            if d.empty:
                continue
            ax.errorbar(d["E_keV"], d[col], yerr=cfg.coverage_k * d[ucol], fmt="o", ms=5, capsize=3, color=PALETTE[i],
                        label=f"{s.name} (deney)")
            if thcol and s.composition and thcol in ("LAC_kuram_cm-1", "MAC_kuram_cm2_g", "HVL_kuram_cm", "Zeff_kuram"):
                comp = ph.normalize_composition(s.composition)
                rho = ph.mean_with_sem(s.density, s.density_u).value if s.density else float("nan")
                mu_rho = ph.mixture_mass_attenuation(xs, comp, egrid)[0]
                if col == "MAC_cm2_g":
                    y = mu_rho
                elif col == "LAC_cm-1":
                    y = mu_rho * rho
                elif col == "HVL_cm":
                    y = ph.LN2 / (mu_rho * rho)
                else:
                    y = np.array([ph.theory_point(xs, comp, e).zeff for e in egrid[::3]])
                xg = egrid if col != "Zeff" else egrid[::3]
                ax.plot(xg, y, color=PALETTE[i], lw=1, alpha=0.7, ls="--")
        ax.set_xscale("log")
        if logy:
            ax.set_yscale("log")
        ax.set_xlabel("Enerji (keV)")
        ax.set_ylabel(label)
        ax.set_title(f"{label}\nnoktalar: deney (±{cfg.coverage_k:g}σ), kesik çizgiler: kuramsal (EPDL97)", fontsize=9)
        ax.legend(fontsize=8)
        ax.grid(alpha=0.3, which="both")
        fig.savefig(outdir / fname, dpi=130)
        plt.close(fig)


# ------------------------------------------------------------------------------ tables
def _summary_table(res: Results) -> pd.DataFrame:
    df, k = res.attenuation, res.cfg.coverage_k
    if df.empty:
        return pd.DataFrame({"Durum": ["Numunesiz (I0) ölçüm tanımlı değil; deneysel parametreler hesaplanamadı."]})
    rows = []
    for _, r in df.iterrows():
        def g(c, uc, nd=0):
            return fmt_pm(r.get(c), r.get(uc), k, nd) if c in r and pd.notna(r.get(c)) else "-"
        rows.append({"Numune": r["Numune"], "Çizgi": r["Cizgi"], "E (keV)": round(r["E_keV"], 1),
                     "T = I/I0": g("T_I_I0", "u_T", 3), "RPE (%)": g("RPE_yuzde", "u_RPE_yuzde"),
                     "LAC μ (cm⁻¹)": g("LAC_cm-1", "u_LAC_cm-1"), "MAC μ/ρ (cm²/g)": g("MAC_cm2_g", "u_MAC_cm2_g"),
                     "HVL (cm)": g("HVL_cm", "u_HVL_cm"), "TVL (cm)": g("TVL_cm", "u_TVL_cm"),
                     "MFP (cm)": g("MFP_cm", "u_MFP_cm"), "Zeff": g("Zeff", "u_Zeff"),
                     "Neff (e/g)": g("Neff_e_g", "u_Neff_e_g"),
                     "Çıkarsanan kalınlık (mm)": g("Kalinlik_cikarsanan_mm", "u_Kalinlik_cikarsanan_mm", 2),
                     "Not": r.get("Not", "") if pd.notna(r.get("Not", "")) else ""})
    return pd.DataFrame(rows)


def _budget_table(res: Results) -> pd.DataFrame:
    df = res.attenuation
    need = {"Katki_ist_I0", "Katki_ist_I", "Katki_model", "Katki_kalinlik"}
    if df.empty or not need.issubset(df.columns):
        return pd.DataFrame()
    d = df.dropna(subset=list(need)).copy()
    var = d[["Katki_ist_I0", "Katki_ist_I", "Katki_model", "Katki_kalinlik"]] ** 2
    tot = var.sum(axis=1)
    out = d[["Numune", "Cizgi", "E_keV", "LAC_cm-1", "u_LAC_cm-1"]].copy()
    out["u: istatistik I0 (cm-1)"] = d["Katki_ist_I0"]
    out["u: istatistik I (cm-1)"] = d["Katki_ist_I"]
    out["u: model/fon secimi (cm-1)"] = d["Katki_model"]
    out["u: kalinlik (cm-1)"] = d["Katki_kalinlik"]
    for c, n in zip(var.columns, ("I0 ist.", "I ist.", "model", "kalinlik")):
        out[f"pay % {n}"] = 100 * var[c] / tot
    return out


METHOD_NOTES = [
    ("Net alan", "Her çizgi grubu (ör. Ba-133 356+384 keV) Gauss + polinom fon (+ Compton kenarı / basamak) modeliyle "
                 "Poisson olabilirlik (sapma) minimizasyonu ile uydurulur. Alan belirsizliği Fisher bilgi matrisinden "
                 "hesaplanır; uyum chi2/ndf>1 ise sqrt(chi2/ndf) (Birge) ile büyütülür."),
    ("Model belirsizliği", "Fon derecesi, uydurma penceresi ve yapısal alternatiflerle (basamak, sigma bağı, Compton kenarı) "
                           "aynı analiz tüm spektrumlara uygulanır. u_model = ln(I0/I) değerlerinin varyantlar arasındaki "
                           "standart sapması (ortak sistematik etkiler oranda otomatik sönümlenir)."),
    ("Sayım hızı", "I = net alan / canlı süre (DSPEC ölü zamanı canlı süre ile düzeltilir)."),
    ("L", "L = ln(I0/I);  u_L² = (u_I0/I0)² + (u_I/I)² + u_model²"),
    ("LAC", "μ = L/x;  u_μ² = (u_L/x)² + (μ u_x/x)²"),
    ("MAC", "μ/ρ = μ/ρ;  u² = (u_μ/ρ)² + (μ u_ρ/ρ²)²  (ρ: deneysel yoğunluk ortalaması ± standart hata)"),
    ("HVL, TVL, MFP", "ln2/μ, ln10/μ, 1/μ;  göreli belirsizlik = u_μ/μ  (yalnızca μ ≥ 3 u_μ ise güvenilir)"),
    ("RPE", "RPE = (1 − I/I0)×100;  u = 100·(I/I0)·u_L"),
    ("Zeff, Neff", "Doğrudan yöntem: σ_a = (μ/ρ)/(N_A Σ w_i/A_i);  σ_e = (1/N_A) Σ f_i (A_i/Z_i)(μ/ρ)_i;  "
                   "Zeff = σ_a/σ_e;  Neff = (μ/ρ)/σ_e;  Ne = Neff·ρ. Elementel (μ/ρ)_i değerleri EPDL97 (XCOM ailesi)."),
    ("Kuramsal", "Karışım kuralı: μ/ρ = Σ w_i (μ/ρ)_i; kuramsal LAC = (μ/ρ)·ρ_deneysel."),
    ("Kapsam", "Tablolar 1σ standart belirsizlik taşır; 'Özet' sayfası yapılandırmadaki kapsam katsayısı k ile verilir."),
]


def write_excel(res: Results, path: Path):
    k = res.cfg.coverage_k
    sheets = {
        "Özet": _summary_table(res),
        "Parametreler": res.attenuation,
        "Pikler": res.peaks,
        "Göreli_zayıflama": res.relative,
        "Teorik": res.theory,
        "Referans_malzeme": res.reference,
        "Numune_bilgisi": res.samples_info,
        "Kalibrasyon": res.calibration,
        "Belirsizlik_bütçesi": _budget_table(res),
        "Uyarılar": pd.DataFrame({"Uyarı": res.warnings or ["(yok)"]}),
        "Yöntem": pd.DataFrame(METHOD_NOTES, columns=["Başlık", "Açıklama"]),
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None or df.empty:
                df = pd.DataFrame({"Bilgi": ["(bu tablo için veri yok - eksik girdi: bkz. Uyarılar)"]})
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            from openpyxl.styles import Alignment, Font, PatternFill
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="305496")
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 70)
            for row in ws.iter_rows(min_row=2):
                for c in row:
                    if isinstance(c.value, float):
                        c.number_format = "0.0000" if abs(c.value) < 1e6 and abs(c.value) >= 1e-3 or c.value == 0 else "0.000E+00"
        ws = xw.sheets["Özet"]
        ws.cell(row=ws.max_row + 2, column=1, value=f"± değerleri: kapsam katsayısı k = {k:g} (k=1: 1σ). "
                                                    "İstatistik, model/fon seçimi ve kalınlık/yoğunluk katkıları dahildir.")


def write_csvs(res: Results, outdir: Path):
    outdir.mkdir(parents=True, exist_ok=True)
    for name, df in {"pikler": res.peaks, "goreli_zayiflama": res.relative, "parametreler": res.attenuation,
                     "teorik": res.theory, "referans_malzeme": res.reference, "numune_bilgisi": res.samples_info,
                     "kalibrasyon": res.calibration, "belirsizlik_butcesi": _budget_table(res)}.items():
        if df is not None and not df.empty:
            df.to_csv(outdir / f"{name}.csv", index=False, encoding="utf-8-sig")


# ------------------------------------------------------------------------------ markdown report
def _md_table(df: pd.DataFrame, cols=None, floatfmt="{:.4g}") -> str:
    if df is None or df.empty:
        return "_(veri yok)_\n"
    d = df[cols] if cols else df
    head = "| " + " | ".join(map(str, d.columns)) + " |\n|" + "|".join(["---"] * len(d.columns)) + "|\n"
    body = ""
    for _, r in d.iterrows():
        cells = []
        for v in r:
            if isinstance(v, (float, np.floating)):
                cells.append("-" if not np.isfinite(v) else floatfmt.format(v))
            else:
                cells.append(str(v))
        body += "| " + " | ".join(cells) + " |\n"
    return head + body


def write_markdown(res: Results, path: Path):
    cfg, st = res.cfg, res.status
    L = [f"# {cfg.project}\n", "## 1. Durum\n"]
    L.append(f"- Spektrum sayısı: {len(res.spectra)} (numune: {len(cfg.samples)}, numunesiz/I0: {'var' if st['blank'] else '**YOK**'})")
    L.append(f"- Kalınlık girilmiş mi: {'evet' if st['thickness_all'] else '**hayır**'}; yoğunluk: {'evet' if st['density_all'] else '**hayır**'}; "
             f"bileşim: {'evet' if st['composition_all'] else '**hayır**'}")
    if not st["blank"] or not st["thickness_all"]:
        L.append("- **LAC, MAC, HVL, TVL, MFP, RPE, Zeff ve Neff henüz hesaplanamadı** (eksik girdi için bkz. Uyarılar).")
    L.append("\n## 2. Veri kontrolü ve kalibrasyon\n")
    L.append(_md_table(res.calibration, ["Spektrum", "Toplam_sayim", "Canli_sure_s", "Gercek_sure_s", "Olu_zaman_yuzde",
                                          "Sayim_hizi_cps", "RMS_artik_keV", "Cozunurluk_FWHM_662_yuzde"]))
    for _, r in res.calibration.iterrows():
        L.append(f"- {r['Spektrum']}: `{r['Kalibrasyon']}`")
    L.append("\n## 3. Tespit edilen çizgiler (net alan, istatistik belirsizlik)\n")
    pk = res.peaks.copy()
    if not pk.empty:
        pk["Alan ± u"] = [fmt_pm(a, u, 1.0) for a, u in zip(pk["Net_alan"], pk["u_alan_ist"])]
        L.append(_md_table(pk[pk["Tespit"] == "evet"], ["Spektrum", "Cizgi", "E_nominal_keV", "Alan ± u", "Anlamlilik_sigma",
                                                        "FWHM_yuzde", "chi2_ndf", "Kabul_varyant"]))
    L.append("\n## 4. Göreli zayıflama (bilgi amaçlı)\n")
    if st["blank"]:
        L.append("Referans: **numunesiz ölçüm (I0)**. T = I/I0; ±: istatistik ⊕ model (kalınlık/yoğunluk belirsizliği hariç).\n")
    else:
        L.append(f"Referans: **{st.get('relative_reference')}**. T = I/I_ref; ±: istatistik ⊕ model. "
                 "(Bu tablo LAC değildir; numunesiz ölçüm yokken spektrumların tutarlılığını gösterir.)\n")
    L.append(_md_table(res.relative, ["Numune", "Cizgi", "E_keV", "T=I/Iref", "u_T", "Anlamlilik_sigma"]))
    if not res.attenuation.empty:
        L.append("\n## 5. Sonuçlar\n")
        L.append(f"Tablo değerleri k = {cfg.coverage_k:g} kapsam katsayısıyla verilmiştir.\n")
        L.append(_md_table(_summary_table(res)))
        ti = res.status.get("thickness_implied")
        if ti:
            L.append("\nVeriden çıkarsanan kalınlık (kuramsal μ ile, tüm çizgilerin ağırlıklı ortalaması): " +
                     "; ".join(f"{k}: {v[0]:.3g} ± {v[1]:.2g} mm" for k, v in ti.items()) +
                     ". Bu değer ölçülen kalınlıkla uyumlu olmalıdır (geometri/numune konumu kontrolü).\n")
    L.append("\n## 6. Kuramsal değerler (karışım kuralı, EPDL97)\n")
    th = res.theory
    L.append(_md_table(th, [c for c in ("Numune", "Cizgi", "E_keV", "MAC_kuram_cm2_g", "LAC_kuram_cm-1", "HVL_kuram_cm",
                                        "Zeff_kuram", "Neff_kuram_e_g") if c in th.columns]) if not th.empty else "")
    L.append("\n## 7. Uyarılar\n")
    L += [f"- {w}" for w in (res.warnings or ["(yok)"])]
    L.append("\n## 8. Yöntem\n")
    L += [f"- **{a}:** {b}" for a, b in METHOD_NOTES]
    path.write_text("\n".join(L) + "\n", encoding="utf-8")


def write_all(res: Results, xs: CrossSections, progress=None) -> Path:
    out = res.cfg.out_dir
    if progress:
        progress("Excel, rapor ve grafikler yazılıyor", 0.92)
    (out / "grafikler").mkdir(parents=True, exist_ok=True)
    write_excel(res, out / "sonuclar.xlsx")
    write_csvs(res, out / "csv")
    write_markdown(res, out / "rapor.md")
    for name in res.analyses:
        safe = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in name)
        plot_spectrum(res, name, out / "grafikler" / f"spektrum_{safe}.png")
    plot_relative(res, out / "grafikler" / "goreli_zayiflama.png")
    plot_parameters(res, xs, out / "grafikler")
    return out
